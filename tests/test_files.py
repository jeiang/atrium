import base64
import io
import json
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import docx as docx_lib
import openpyxl
import pytest
from fake_anthropic import FakeClient, response, text, tool_use
from helpers import make_config, static_dir, web_client

from atrium import db
from atrium.agent import files, journal, loop, projects, registry, service, threads, workspace
from atrium.agent.tools.base import Tool

HX = {"HX-Request": "true"}
NOW = 1_790_000_000.0
PDF = b"%PDF-1.4\n1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n2 0 obj << /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >> endobj\n3 0 obj << /Type /Page >> endobj\n4 0 obj << /Type /Page >> endobj\n%%EOF"
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def probe(ctx, args):
    return {"echo": args.get("x")}


TOOLS = [Tool("probe", "probe", {"type": "object"}, probe)]


@pytest.fixture
def cfg(tmp_path, monkeypatch):
    monkeypatch.setattr(registry, "tools_for", lambda kind: [t for t in TOOLS if kind in t.contexts])
    return make_config(tmp_path, static_dir=static_dir(tmp_path))


@pytest.fixture
def conn(cfg):
    c = db.connect(cfg.db_path)
    yield c
    c.close()


class FakeFiles:
    def __init__(self):
        self.uploaded, self.deleted, self.store = [], [], {}

    def upload(self, file):
        name, data, media_type = file
        file_id = f"file_{len(self.uploaded) + 1}"
        self.uploaded.append((name, media_type))
        self.store[file_id] = (name, data, media_type)
        return SimpleNamespace(id=file_id)

    def delete(self, file_id):
        self.deleted.append(file_id)

    def retrieve_metadata(self, file_id):
        name, data, media_type = self.store[file_id]
        return SimpleNamespace(id=file_id, filename=name, mime_type=media_type, size_bytes=len(data))

    def download(self, file_id):
        return SimpleNamespace(read=lambda: self.store[file_id][1])


class SandboxClient(FakeClient):
    def __init__(self, responses):
        super().__init__(responses)
        self.files = FakeFiles()


def code_execution(use_id, command, stdout, file_ids=(), return_code=0):
    return [
        {"type": "server_tool_use", "id": use_id, "name": "bash_code_execution", "input": {"command": command}},
        {
            "type": "bash_code_execution_tool_result",
            "tool_use_id": use_id,
            "content": {
                "type": "bash_code_execution_result", "stdout": stdout, "stderr": "", "return_code": return_code,
                "content": [{"type": "bash_code_execution_output", "file_id": f} for f in file_ids],
            },
        },
    ]


def docx_bytes(*paragraphs):
    document = docx_lib.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "qty", "7"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx_bytes():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(["month", "total"])
    sheet.append(["Jan", 120])
    other = book.create_sheet("Notes")
    other.append(["remember, commas"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def upload(web, thread_id, text_in, *parts, edit=""):
    return web.post(
        f"/chat/{thread_id}/messages", data={"text": text_in, "edit": edit},
        files=[("upload", (name, data, media)) for name, data, media in parts], headers=HX,
    )


def wait_done(cfg, thread_id, timeout=5.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        other = db.connect(cfg.db_path)
        try:
            runs = journal.runs_of(other, thread_id)
        finally:
            other.close()
        if runs and runs[-1]["state"] in journal.TERMINAL:
            return runs[-1]
        time.sleep(0.05)
    raise AssertionError("run did not finish")


def new_thread(web, text_in="", *parts, project_id=None):
    data = {"text": text_in}
    if project_id:
        data["project_id"] = str(project_id)
    reply = web.post("/chat", data=data, files=[("upload", (n, d, m)) for n, d, m in parts], headers=HX)
    assert reply.status_code == 204, reply.text
    return int(reply.headers["hx-redirect"].rsplit("/", 1)[1])


def last_request_user_blocks(client):
    return client.calls[-1]["messages"][-1]["content"]


# ---- uploads: storage, listing, previews, downloads


def test_upload_is_stored_under_the_state_dir_and_listed_in_work(cfg, conn):
    client = FakeClient([response([text("got it")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(web, "look at this", ("data.csv", b"a,b\n1,2\n", "text/csv"))
    wait_done(cfg, thread_id)
    (row,) = workspace.files_of(conn, thread_id)
    assert row["origin"] == "upload" and row["media_type"] == "text/csv" and row["bytes"] == 8 and row["text"] == "a,b\n1,2\n"
    path = Path(row["path"])
    assert path.is_file() and path.parent == files.files_dir(cfg) / str(thread_id) and path.name == f"{row['id']}-data.csv"
    assert files.files_dir(cfg).parent == cfg.db_path.parent
    page = web.get(f"/chat/{thread_id}").text
    assert "data.csv" in page and f'hx-get="/chat/{thread_id}/files/{row["id"]}"' in page and f"/files/{row['id']}/download" in page and "data-attachments" in page
    message = journal.messages_of(conn, thread_id)[0]
    assert [b["type"] for b in message["content"]] == ["attachment", "text"]
    assert message["content"][0] == {"type": "attachment", "file_id": row["id"], "name": "data.csv", "media_type": "text/csv", "kind": "text", "tokens": 2}
    view = service.thread_view(conn, thread_id)
    assert view["items"][0]["files"][0]["name"] == "data.csv"


def test_each_type_becomes_the_right_content_block(cfg, conn):
    client = FakeClient([response([text("read them")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(
        web, "summarize",
        ("report.pdf", PDF, "application/pdf"), ("pic.png", PNG, "image/png"), ("notes.md", b"# hi\nsee https://evil.test/x", "text/markdown"),
        ("brief.docx", docx_bytes("Hello from Word", "Second paragraph"), "application/octet-stream"), ("sales.xlsx", xlsx_bytes(), "application/octet-stream"),
    )
    wait_done(cfg, thread_id)
    blocks = last_request_user_blocks(client)
    kinds = [(b["type"], b.get("source", {}).get("type"), b.get("source", {}).get("media_type")) for b in blocks]
    assert kinds[:5] == [
        ("document", "base64", "application/pdf"), ("image", "base64", "image/png"), ("document", "text", "text/plain"),
        ("document", "text", "text/plain"), ("document", "text", "text/plain"),
    ]
    assert blocks[5]["type"] == "text" and blocks[5]["text"] == "summarize" and blocks[-1]["text"].startswith("Current date")
    assert base64.b64decode(blocks[0]["source"]["data"]) == PDF and blocks[0]["citations"] == {"enabled": True} and blocks[0]["title"] == "report.pdf"
    assert "untrusted" in blocks[0]["context"] and "untrusted" in blocks[2]["context"]
    assert base64.b64decode(blocks[1]["source"]["data"]) == PNG
    assert blocks[2]["source"]["data"] == "# hi\nsee https://evil.test/x"
    assert "Hello from Word" in blocks[3]["source"]["data"] and "Second paragraph" in blocks[3]["source"]["data"] and "qty\t7" in blocks[3]["source"]["data"]
    assert blocks[4]["source"]["data"].startswith("# Sheet: Sales\nmonth,total\nJan,120") and '"remember, commas"' in blocks[4]["source"]["data"]
    assert "cache_control" in blocks[-1] and sum("cache_control" in b for b in blocks) == 1
    assert threads.allowed_urls(conn, thread_id) == set()  # a URL inside a file grants nothing
    rows = workspace.files_of(conn, thread_id)
    assert [r["media_type"] for r in rows][3:] == [files.DOCX, files.XLSX]
    assert rows[0]["tokens"] == 2 * files.PDF_TOKENS_PER_PAGE and rows[1]["tokens"] == files.IMAGE_TOKENS
    assert "container_upload" not in {b["type"] for b in blocks} and "tools" in client.calls[0]
    assert all(t.get("type") != files.CODE_EXECUTION_TOOL["type"] for t in client.calls[0]["tools"])


def test_attachments_persist_across_turns_and_are_rebuilt_each_request(cfg, conn):
    client = FakeClient([response([text("one")]), response([text("two")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(web, "first", ("report.pdf", PDF, "application/pdf"))
    wait_done(cfg, thread_id)
    assert upload(web, thread_id, "and again").status_code == 200
    wait_done(cfg, thread_id)
    second = client.calls[1]["messages"]
    assert second[0]["content"][0]["type"] == "document" and base64.b64decode(second[0]["content"][0]["source"]["data"]) == PDF
    assert second[-1]["content"][0]["text"] == "and again"
    stored = journal.messages_of(conn, thread_id)
    assert all("base64" not in json.dumps(m["content"]) for m in stored)


def test_oversize_and_unreadable_files_are_refused(cfg, conn, monkeypatch):
    client = FakeClient([response([text("fits now")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = service.create_thread(conn, "chat")
    monkeypatch.setattr(files, "FILE_CAP_BYTES", 10)
    refused = upload(web, thread_id, "big", ("big.txt", b"x" * 11, "text/plain"))
    assert refused.status_code == 422 and "over the 0 MB limit" in refused.text
    monkeypatch.setattr(files, "FILE_CAP_BYTES", 32 << 20)
    binary = upload(web, thread_id, "bin", ("blob.bin", b"\0\1\2", "application/octet-stream"))
    assert binary.status_code == 422 and "atrium reads" in binary.text
    huge = upload(web, thread_id, "huge", ("huge.txt", b"w" * 400_000, "text/plain"))
    assert huge.status_code == 422 and "context" in huge.text and "Raise the context limit" in huge.text
    assert workspace.files_of(conn, thread_id) == [] and journal.messages_of(conn, thread_id) == [] and client.calls == []
    assert not (files.files_dir(cfg) / str(thread_id)).exists()
    threads.set_options(conn, thread_id, "medium", 1_000_000)
    assert upload(web, thread_id, "huge", ("huge.txt", b"w" * 400_000, "text/plain")).status_code == 200
    wait_done(cfg, thread_id)
    assert len(client.calls) == 1 and client.calls[0]["messages"][0]["content"][0]["source"]["data"] == "w" * 400_000


def test_files_without_text_and_text_without_files(cfg, conn):
    web = web_client(cfg, agent_client=lambda: FakeClient([]))
    reply = web.post("/chat", data={"text": ""}, files=[("upload", ("a.txt", b"hi", "text/plain"))], headers=HX)
    assert reply.status_code == 422 and "Type a message" in reply.text
    thread_id = service.create_thread(conn, "chat")
    assert upload(web, thread_id, "").status_code == 422


def test_preview_and_download_headers(cfg, conn):
    client = FakeClient([response([text("ok")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(
        web, "files", ("page.html", b"<script>alert(1)</script>", "text/html"), ("pic.png", PNG, "image/png"),
        ("brief.docx", docx_bytes("Docx body"), files.DOCX), ("sales.xlsx", xlsx_bytes(), files.XLSX),
    )
    wait_done(cfg, thread_id)
    html, png, word, sheet = workspace.files_of(conn, thread_id)
    got = web.get(f"/chat/{thread_id}/files/{html['id']}/download")
    assert got.status_code == 200 and got.headers["content-disposition"] == 'attachment; filename="page.html"'
    assert got.headers["content-type"].startswith("text/html") and got.headers["x-content-type-options"] == "nosniff" and got.content == b"<script>alert(1)</script>"
    assert web.get(f"/chat/{thread_id}/files/{html['id']}/raw").status_code == 404
    inline = web.get(f"/chat/{thread_id}/files/{png['id']}/raw")
    assert inline.status_code == 200 and inline.headers["content-type"] == "image/png" and inline.content == PNG
    preview = web.get(f"/chat/{thread_id}/files/{png['id']}", headers=HX).text
    assert f'<img class="max-w-full rounded-md border" src="/chat/{thread_id}/files/{png["id"]}/raw"' in preview and f"/files/{png['id']}/download" in preview
    assert "Docx body" in web.get(f"/chat/{thread_id}/files/{word['id']}", headers=HX).text
    table = web.get(f"/chat/{thread_id}/files/{sheet['id']}", headers=HX).text
    assert "<table" in table and "<td" in table and "Jan" in table
    escaped = web.get(f"/chat/{thread_id}/files/{html['id']}", headers=HX).text
    assert "<script>alert" not in escaped and "&lt;script&gt;" in escaped
    assert web.get(f"/chat/{thread_id}/files/{html['id'] + 99}/download").status_code == 404
    other = service.create_thread(conn, "chat")
    assert web.get(f"/chat/{other}/files/{html['id']}/download").status_code == 404


def test_queued_message_carries_its_attachments(cfg, conn):
    release = threading.Event()

    class Slow(FakeClient):
        def create(self, **params):
            release.wait(5)
            return super().create(**params)

    client = Slow([response([tool_use("tu1", "probe", x=1)], "tool_use"), response([text("done")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(web, "start")
    queued = upload(web, thread_id, "also this", ("extra.txt", b"more words", "text/plain"))
    assert queued.status_code == 200 and "Queued" in queued.text and "extra.txt" in queued.text
    release.set()
    wait_done(cfg, thread_id)
    second = client.calls[1]["messages"][-1]["content"]
    assert [b["type"] for b in second][:3] == ["tool_result", "document", "text"] and second[1]["source"]["data"] == "more words"


# ---- code sandbox


def test_sandbox_off_sends_no_code_execution_tool_and_touches_no_files_api(cfg, conn):
    client = FakeClient([response([text("plain")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(web, "hello", ("data.csv", b"a,b\n", "text/csv"))
    wait_done(cfg, thread_id)
    call = client.calls[0]
    assert all(t.get("type") != files.CODE_EXECUTION_TOOL["type"] for t in call["tools"]) and "container" not in call
    assert all(b["type"] != "container_upload" for b in call["messages"][0]["content"])
    assert not any(files.SANDBOX_PROMPT in b["text"] for b in call["system"])
    assert workspace.files_of(conn, thread_id)[0]["anthropic_file_id"] is None
    page = web.get(f"/chat/{thread_id}").text
    assert 'name="code_sandbox"' in page and 'id="sandbox" class="flex items-center gap-1 text-xs text-warning" role="status" hidden' in page


def test_sandbox_on_uploads_to_files_api_and_enables_the_tool(cfg, conn):
    client = SandboxClient([response([text("ready")]), response([text("again")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(web, "first")
    wait_done(cfg, thread_id)
    toggled = web.post(f"/chat/{thread_id}/sandbox", data={"code_sandbox": "on"}, headers=HX)
    assert toggled.status_code == 200 and "not zero data retention" in toggled.text and "30 days" in toggled.text and 'role="status" hidden' not in toggled.text
    assert threads.get(conn, thread_id)["code_sandbox"] == 1
    assert upload(web, thread_id, "analyze", ("data.csv", b"a,b\n1,2\n", "text/csv")).status_code == 200
    wait_done(cfg, thread_id)
    assert client.files.uploaded == [("data.csv", "text/csv")]
    row = workspace.files_of(conn, thread_id)[0]
    assert row["anthropic_file_id"] == "file_1"
    call = client.calls[1]
    assert call["tools"][-1] == {**files.CODE_EXECUTION_TOOL, "cache_control": {"type": "ephemeral"}}
    assert any(files.SANDBOX_PROMPT in b["text"] for b in call["system"])
    blocks = call["messages"][-1]["content"]
    assert [b["type"] for b in blocks][:3] == ["document", "container_upload", "text"] and blocks[1]["file_id"] == "file_1"
    off = web.post(f"/chat/{thread_id}/sandbox", data={}, headers=HX)
    assert off.status_code == 200 and 'role="status" hidden' in off.text and threads.get(conn, thread_id)["code_sandbox"] == 0
    assert web.post(f"/chat/{thread_id}/sandbox", data={"code_sandbox": "on"}, follow_redirects=False).status_code == 303


def test_generated_files_are_fetched_and_shown_as_code_sessions(cfg, conn):
    client = SandboxClient(
        [
            response(code_execution("srv1", "python plot.py && cp chart.png $OUTPUT_DIR/", "saved\n", ["gen_1"]) + [tool_use("tu1", "probe", x=1)], "tool_use"),
            response([text("Here is the chart.")]),
        ]
    )
    client.files.store["gen_1"] = ("chart.png", PNG, "image/png")
    thread_id = service.create_thread(conn, "chat")
    threads.set_flag(conn, thread_id, "code_sandbox", True)
    run_id = service.run_once(cfg, conn, thread_id, "plot it", kind="chat", client=client, env={})
    assert journal.get_run(conn, run_id)["state"] == journal.DONE
    (made,) = workspace.files_of(conn, thread_id)
    assert made["origin"] == "generated" and made["name"] == "chart.png" and made["media_type"] == "image/png" and made["anthropic_file_id"] == "gen_1"
    assert Path(made["path"]).read_bytes() == PNG and made["run_id"] == run_id
    (session,) = workspace.code_sessions(conn, thread_id)
    assert session["tool"] == "code_execution" and session["script"].startswith("python plot.py") and session["result"] == "saved\n" and session["state"] == "done"
    items = service.thread_view(conn, thread_id)["items"]
    step = next(s for s in items[1]["steps"] if s["kind"] == "tool")
    assert step["name"] == "bash_code_execution" and step["result"] == "saved\n" and not step["is_error"]
    page = web_client(cfg).get(f"/chat/{thread_id}").text
    assert "chart.png" in page and "generated" in page and "code_execution" in page and "data-code" in page
    again = FakeClient([response(code_execution("srv2", "ls", "x", ["gen_1"]), "end_turn")])
    again.files = client.files
    service.run_once(cfg, conn, thread_id, "once more", kind="chat", client=again, env={})
    assert len(workspace.files_of(conn, thread_id)) == 1  # the same remote file is not stored twice


def test_container_is_reused_and_replaced_when_rejected(cfg, conn):
    class WithContainer(SandboxClient):
        def create(self, **params):
            found = super().create(**params)
            found.container = SimpleNamespace(id="cont_1", expires_at=None)
            return found

    client = WithContainer([response([text("a")]), response([text("b")])])
    thread_id = service.create_thread(conn, "chat")
    threads.set_flag(conn, thread_id, "code_sandbox", True)
    service.run_once(cfg, conn, thread_id, "one", kind="chat", client=client, env={})
    assert threads.get(conn, thread_id)["container_id"] == "cont_1" and "container" not in client.calls[0]
    service.run_once(cfg, conn, thread_id, "two", kind="chat", client=client, env={})
    assert client.calls[1]["container"] == "cont_1"

    class Expired(SandboxClient):
        def create(self, **params):
            if "container" in params:
                self.calls.append(params)
                raise type("BadRequest", (Exception,), {"status_code": 400})("container expired")
            return super().create(**params)

    retry = Expired([response([text("fresh")])])
    service.run_once(cfg, conn, thread_id, "three", kind="chat", client=retry, env={})
    assert [("container" in c) for c in retry.calls] == [True, False] and threads.get(conn, thread_id)["container_id"] is None


# ---- deletion


def test_deleting_a_thread_removes_files_locally_and_remotely(cfg, conn):
    client = SandboxClient([response([text("ok")]), response([text("ok")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(web, "hello")
    wait_done(cfg, thread_id)
    web.post(f"/chat/{thread_id}/sandbox", data={"code_sandbox": "on"}, headers=HX)
    upload(web, thread_id, "take this", ("data.csv", b"a,b\n", "text/csv"))
    wait_done(cfg, thread_id)
    folder = files.files_dir(cfg) / str(thread_id)
    assert folder.is_dir() and client.files.uploaded
    assert thread_id in {t["id"] for t in service.list_threads(conn)}
    reply = web.post(f"/chat/{thread_id}/delete", follow_redirects=False)
    assert reply.status_code == 303 and reply.headers["location"] == "/chat"
    assert not folder.exists() and workspace.files_of(conn, thread_id) == [] and client.files.deleted == ["file_1"]
    assert journal.messages_of(conn, thread_id) == [] and threads.get(conn, thread_id) is None
    assert thread_id not in {t["id"] for t in service.list_threads(conn)} and thread_id not in {t["id"] for t in service.list_threads(conn, archived=True)}
    assert web.get(f"/chat/{thread_id}").status_code == 404 and web.post(f"/chat/{thread_id}/delete").status_code == 404
    assert conn.execute("SELECT COUNT(*) FROM usage WHERE thread_id = ?", (thread_id,)).fetchone()[0] == 2  # spend accounting survives


def test_branch_copies_attachments_so_either_thread_can_be_deleted(cfg, conn):
    client = FakeClient([response([text("one")]), response([text("two")])])
    web = web_client(cfg, agent_client=lambda: client)
    thread_id = new_thread(web, "first", ("notes.txt", b"keep me", "text/plain"))
    wait_done(cfg, thread_id)
    seq = journal.messages_of(conn, thread_id)[0]["seq"]
    child = service.branch(conn, thread_id, seq)
    (copy,) = workspace.files_of(conn, child)
    original = workspace.files_of(conn, thread_id)[0]
    assert copy["id"] != original["id"] and Path(copy["path"]).read_bytes() == b"keep me" and copy["path"] != original["path"]
    assert journal.messages_of(conn, child)[0]["content"][0]["file_id"] == copy["id"]
    service.delete_thread(cfg, conn, thread_id, None)
    assert Path(copy["path"]).is_file()
    service.run_once(cfg, conn, child, "go on", kind="chat", client=client, env={})
    assert client.calls[-1]["messages"][0]["content"][0]["source"]["data"] == "keep me"


# ---- project files


def test_project_files_are_attached_to_every_request_of_its_threads(cfg, conn):
    project_id = projects.create(conn, NOW, name="Taxes", instructions="", effort="medium", context_limit=90_000)
    client = FakeClient([response([text("sure")]), response([text("more")])])
    web = web_client(cfg, agent_client=lambda: client)
    listed = web.post(f"/projects/{project_id}/files", files=[("upload", ("rules.md", b"# Rules\nfile early", "text/markdown"))], headers=HX)
    assert listed.status_code == 200 and "rules.md" in listed.text and f"/projects/{project_id}/files/" in listed.text
    (row,) = files.of_project(conn, project_id)
    assert row["thread_id"] is None and row["project_id"] == project_id and Path(row["path"]).parent == files.files_dir(cfg) / f"p{project_id}"
    page = web.get(f"/projects/{project_id}").text
    assert "rules.md" in page and "data-project-files" in page
    got = web.get(f"/projects/{project_id}/files/{row['id']}/download")
    assert got.status_code == 200 and got.headers["content-disposition"] == 'attachment; filename="rules.md"'
    thread_id = new_thread(web, "what should I do?", project_id=project_id)
    wait_done(cfg, thread_id)
    first = client.calls[0]["messages"][0]["content"]
    assert first[0]["type"] == "document" and first[0]["source"]["data"] == "# Rules\nfile early" and first[0]["title"] == "rules.md"
    assert "Project files" in first[1]["text"] and first[2]["text"] == "what should I do?"
    assert workspace.files_of(conn, thread_id) == []  # the project's files are not the thread's
    assert upload(web, thread_id, "next").status_code == 200
    wait_done(cfg, thread_id)
    second = client.calls[1]["messages"]
    assert second[0]["content"][0]["source"]["data"] == "# Rules\nfile early" and second[-1]["content"][0]["text"] == "next"
    rejected = web.post(f"/projects/{project_id}/files", files=[("upload", ("blob.bin", b"\0\0", "application/octet-stream"))], headers=HX)
    assert rejected.status_code == 422 and "atrium reads" in rejected.text
    removed = web.post(f"/projects/{project_id}/files/{row['id']}/delete", headers=HX)
    assert removed.status_code == 200 and "rules.md" not in removed.text and files.of_project(conn, project_id) == [] and not Path(row["path"]).exists()
    assert web.post(f"/projects/{project_id}/files/{row['id']}/delete", headers=HX).status_code == 404


def test_loop_estimate_counts_hidden_pdf_tokens(cfg, conn):
    thread_id = service.create_thread(conn, "chat")
    threads.set_options(conn, thread_id, "medium", 90_000)
    block = {"type": "attachment", "file_id": 1, "name": "big.pdf", "media_type": "application/pdf", "kind": "pdf", "tokens": 95_000}
    active = [{"role": "user", "content": [block, {"type": "text", "text": "hi"}], "seq": 1, "run_id": None}]
    assert loop.estimate({"system": [], "tools": []}, active, None) > 95_000
