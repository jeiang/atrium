"""Anthropic request goldens: every Messages request the frozen Python loop sends, captured on the wire.

The loop runs against tests/fake_anthropic.py for its responses; each call is also replayed through the real
anthropic SDK onto a transport that records the HTTP request and fails it, so the recorded body is the one the
SDK serializes. Output: testdata/golden/anthropic/ (or argv[1])."""

import base64
import io
import json
import os
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))

from common import FREEZE_COMMIT, ROOT, dumps, out_dir, use_source_tree, write_json  # noqa: E402

TZ = "America/Port_of_Spain"
os.environ["ATRIUM_TZ"] = TZ
for key in ("ANTHROPIC_BASE_URL", "ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN"):
    os.environ.pop(key, None)
use_source_tree()

import anthropic  # noqa: E402
import docx  # noqa: E402
import httpx2  # noqa: E402
import openpyxl  # noqa: E402

from atrium import db  # noqa: E402
from atrium.agent import budget, files, loop, projects, registry, scheduler, service, skills, subagents, threads  # noqa: E402
from atrium.agent.nightly import NIGHTLY_PROMPT  # noqa: E402
from atrium.agent.tools import all_tools, run_code  # noqa: E402
from atrium.agent.tools.base import CHAT, CRON, NIGHTLY, SUB  # noqa: E402
from atrium.agent.tools.profile import profile_block  # noqa: E402
from atrium.config import Config  # noqa: E402
from fake_anthropic import FakeClient, FakeStream, cleared, compaction, iteration, response, text, thinking, tool_use, web_search  # noqa: E402

SLICE = "anthropic"
ZONE = ZoneInfo(TZ)
NOW = datetime(2026, 10, 15, 12, 0, tzinfo=ZONE).timestamp()
LATER = datetime(2026, 11, 2, 9, 30, tzinfo=ZONE).timestamp()
BASE_URL = "https://api.anthropic.com"
HEADERS = ("anthropic-beta", "anthropic-version", "content-type")
HUB = "0f1e2d3c4b5a69788796a5b4c3d2e1f0"
PROFILE = (("city", "Port of Spain"), ("tone", "brief, no emoji"))
EXCLUDED = {
    "tools": ["run_code"],
    "paths": ["code_execution tool and container (thread code sandbox)", "container_upload blocks", "Files API upload and delete", "notion_migrate json_schema output_config.format"],
    "reason": "not ported in Phase A (https://github.com/jeiang/atrium/issues/19, https://github.com/jeiang/atrium/issues/96 review fixes S7)",
}
DIVERGENCES = [
    {
        "where": "prompts.SYSTEM (system[0].text of every request here)",
        "what": "tells the model to call run_code for several read-only lookups; run_code is not ported in Phase A and is absent from every tools list here",
        "owner": "the Go prompt slice (https://github.com/jeiang/atrium/issues/125) decides how to handle the text; until it does, a Go body may differ from these goldens at system[0].text only by the run_code sentence",
    },
    {
        "where": "files.SANDBOX_PROMPT (sent only with the thread code sandbox on, so in no golden here)",
        "what": "names run_code as a separate interpreter; the code sandbox and run_code are both not ported in Phase A",
        "owner": "the Go prompt slice (https://github.com/jeiang/atrium/issues/125)",
    },
]
PDF = b"%PDF-1.4\n1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n2 0 obj << /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >> endobj\n3 0 obj << /Type /Page >> endobj\n4 0 obj << /Type /Page >> endobj\n%%EOF"
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


class Captured(Exception):
    pass


def plain(value):
    if isinstance(value, SimpleNamespace):
        return {k: plain(v) for k, v in vars(value).items()}
    if isinstance(value, dict):
        return {k: plain(v) for k, v in value.items()}
    if isinstance(value, list | tuple):
        return [plain(v) for v in value]
    return value


class Recorder(FakeClient):
    """FakeClient whose every call is also sent through the real SDK to capture the wire request."""

    def __init__(self, responses):
        super().__init__(responses)
        self.exchanges = []
        self.sdk = anthropic.Anthropic(api_key="golden", base_url=BASE_URL, max_retries=0, http_client=httpx2.Client(transport=httpx2.MockTransport(self.capture)))

    def capture(self, request):
        url = request.url
        self.exchanges.append(
            {
                "request": {
                    "method": request.method,
                    "path": url.path,
                    "query": url.query.decode() if isinstance(url.query, bytes) else str(url.query),
                    "headers": {k: request.headers[k] for k in HEADERS if k in request.headers},
                    "body": json.loads(request.content),
                }
            }
        )
        raise Captured

    def wire(self, call):
        try:
            call()
        except anthropic.APIConnectionError as e:
            if not isinstance(e.__cause__, Captured):
                raise
        else:
            raise AssertionError("request was not captured")

    def reply(self, params):
        found = FakeClient.create(self, **params)
        self.exchanges[-1]["response"] = plain(found)
        return found

    def create(self, **params):
        self.wire(lambda: self.sdk.beta.messages.create(**params))
        return self.reply(params)

    def stream(self, **params):
        def send():
            with self.sdk.beta.messages.stream(**params):
                pass

        self.wire(send)
        return FakeStream(self.reply(params))


def ported_tools():
    return [t for t in all_tools() if t.name != run_code.TOOL_NAME]


def clock(start):
    state = {"t": start}

    def now():
        state["t"] += 1
        return state["t"]

    return now


class World:
    def __init__(self, root):
        self.cfg = Config(
            db_path=root / "atrium.db",
            cache_dir=root / "cache",
            mode="shadow",
            chat_url="http://chat.invalid/v1",
            embed_url="http://embed.invalid/v1",
            contacts_dir=None,
            credentials={},
        )
        self.conn = db.connect(self.cfg.db_path)
        for key, value in PROFILE:
            self.conn.execute("INSERT INTO agent_profile (key, value, updated_at) VALUES (?, ?, ?)", (key, value, NOW - 86400))

    def chat(self, client, message, thread_id=None, kind=CHAT, now=None, **thread):
        now = now or clock(NOW)
        if thread_id is None:
            thread_id = service.create_thread(self.conn, kind, now=now, **thread)
        service.run_once(self.cfg, self.conn, thread_id, message, kind=kind, client=client, env={}, now=now, operator=kind == CHAT)
        return thread_id

    def chat_with_files(self, client, message, uploads, thread_id=None, now=None, **thread):
        """What the web chat form does: ingest the uploads, journal them ahead of the text, run."""
        now = now or clock(NOW)
        if thread_id is None:
            thread_id = service.create_thread(self.conn, CHAT, now=now, **thread)
        found = threads.get(self.conn, thread_id)
        blocks = files.ingest_uploads(self.cfg, self.conn, uploads, now(), thread_id=thread_id, context_limit=threads.context_limit_of(found))
        run_id = service.prepare(self.conn, thread_id, message, CHAT, now(), operator=True, files=blocks)
        service.execute(self.cfg, self.conn, run_id, client, {}, now)
        return thread_id

    def job(self, client, kind, due_at):
        claims, skipped = scheduler.claim_due(self.conn, self.cfg, due_at + 5, "golden")
        assert len(claims) == 1 and not skipped and claims[0].job["kind"] == kind, claims
        status = scheduler.execute_claim(self.cfg, claims[0], client=client, env={}, now=clock(due_at + 5))
        assert status == scheduler.OK, status

    def records(self):
        runs = [dict(r) for r in self.conn.execute("SELECT id, thread_id, parent_run_id, kind, effort, state, error, steps FROM runs ORDER BY id")]
        usage_rows = [dict(r) for r in self.conn.execute("SELECT * FROM usage ORDER BY id")]
        for row in usage_rows:
            row.pop("id")
        return runs, usage_rows


SCENARIOS = []


def scenario(name, description, source_tests):
    def register(fn):
        SCENARIOS.append((name, description, source_tests, fn))
        return fn

    return register


@scenario(
    "chat_tool_turn",
    "Chat run at the default effort and 5m TTL: tool_use (profile_get) then end_turn. System = SYSTEM + profile + skills index; breakpoints on SYSTEM, the last stable block, the last tool and the last user block.",
    [
        "tests/test_agent_loop.py::test_loop_executes_tools_journals_intent_and_returns_thinking_unchanged",
        "tests/test_agent_loop.py::test_request_shape_pins_effort_caching_and_stamp",
        "tests/test_agent_loop.py::test_profile_block_is_a_separate_cached_system_block",
        "tests/test_agent_context.py::test_request_carries_context_management_and_betas",
    ],
)
def chat_tool_turn(world):
    client = Recorder(
        [
            response([thinking("check the profile", "sig-1"), tool_use("toolu_01", "profile_get")], "tool_use"),
            response([text("You are in Port of Spain.")]),
        ]
    )
    world.chat(client, "Where do I live? See https://example.test/a")
    return client


@scenario(
    "chat_followup_turn",
    "Two chat runs in one thread: each operator message keeps the date stamp of its own run and the history breakpoint moves to the newest user block.",
    ["tests/test_agent_context.py::test_breakpoints_stay_within_four_and_stable_across_turns"],
)
def chat_followup_turn(world):
    client = Recorder([response([text("Hello.")]), response([text("Still here.")])])
    thread_id = world.chat(client, "hello")
    world.chat(client, "are you there?", thread_id=thread_id, now=clock(LATER))
    return client


@scenario(
    "cache_ttl_1h",
    "cache_ttl_chat = 1h: the chat run's prefix and history breakpoints carry ttl 1h (bug #67); a cron run in the same instance keeps the 5m default.",
    ["tests/test_agent_context.py::test_chat_prefix_ttl_setting"],
)
def cache_ttl_1h(world):
    budget.set_setting(world.conn, loop.CACHE_TTL_KEY, "1h")
    client = Recorder(
        [
            response([tool_use("toolu_01", "context_status")], "tool_use"),
            response([text("Plenty of room.")]),
            response([text("tick done")]),
        ]
    )
    world.chat(client, "how full is the context?")
    world.chat(client, "tick", kind=CRON)
    return client


@scenario(
    "pause_turn",
    "A response stopped with pause_turn (server tool still running) is journaled and the loop sends the continuation request unchanged in shape.",
    ["src/atrium/agent/loop.py::run (PAUSE_TURN branch; no dedicated test)"],
)
def pause_turn(world):
    searched = web_search("srvtoolu_01", "carnival dates 2027", ["https://carnival.test/2027"])
    client = Recorder(
        [
            response([thinking("search first", "sig-p"), *searched], "pause_turn", searches=1),
            response([text("Carnival 2027 is on 8-9 February [https://carnival.test/2027].")]),
        ]
    )
    world.chat(client, "When is carnival in 2027?")
    return client


@scenario(
    "web_search",
    "web_search_20260318 declared last with allowed_callers [direct] (bug #68); the server_tool_use and web_search_tool_result blocks return unchanged in the next turn.",
    [
        "tests/test_agent_loop.py::test_web_search_is_direct_only_in_request_body",
        "tests/test_agent_loop.py::test_web_search_urls_become_allowed",
    ],
)
def web_search_turns(world):
    client = Recorder(
        [
            response([*web_search("srvtoolu_01", "atrium release notes", ["https://news.test/a", "https://news.test/b"]), text("Two articles found.")], searches=1),
            response([text("The first article [https://news.test/a] covers the 2.0 release.")]),
        ]
    )
    thread_id = world.chat(client, "search the web for atrium release notes")
    world.chat(client, "open the first one", thread_id=thread_id, now=clock(LATER))
    return client


@scenario(
    "server_compaction",
    "Server-side compaction (compact_20260112, pause_after_compaction): the compaction block opens the next request, followed by the context manifest and the operator's message with its stamp; thinking is stripped.",
    ["tests/test_agent_context.py::test_server_compaction_journals_block_and_continues_with_manifest_and_stripped_thinking"],
)
def server_compaction(world):
    client = Recorder(
        [
            response([thinking("deep", "sig-1"), tool_use("toolu_01", "profile_get")], "tool_use"),
            response([compaction("Goal: answer where the operator lives. Pending: none.")], "compaction", input_tokens=80_000),
            response([text("Port of Spain.")]),
        ]
    )
    world.chat(client, "where do I live?")
    return client


@scenario(
    "usage_iterations",
    "A compaction response whose usage.iterations bills every iteration (bug #65); the usage rows hold the summed tokens and the >100K tier.",
    ["tests/test_agent_context.py::test_server_compaction_bills_every_usage_iteration"],
)
def usage_iterations(world):
    iterations = [iteration("compaction", input_tokens=300_000, output_tokens=4_000), iteration("message", input_tokens=2_000, cache_read=10_000, output_tokens=50)]
    client = Recorder(
        [
            response([compaction("Goal: answer hello.")], "compaction", input_tokens=2_000, cache_read=10_000, output_tokens=50, iterations=iterations),
            response([text("done")]),
        ]
    )
    world.chat(client, "hello")
    return client


@scenario(
    "client_compaction",
    "Preflight over 95% of the 90K limit compacts on the client: one tool-free, thinking-free create() at effort low with the summary system prompt, then the run resumes from summary + manifest + the operator's message.",
    [
        "tests/test_agent_context.py::test_preflight_compacts_on_the_client_before_exceeding_the_limit",
        "tests/test_agent_context.py::test_client_compaction_mid_tool_loop_sends_a_valid_tool_free_request",
    ],
)
def client_compaction(world):
    client = Recorder(
        [
            response([thinking("t", "s"), tool_use("toolu_01", "profile_get")], "tool_use", input_tokens=90_000, output_tokens=500),
            response([text("SUMMARY: goal hello; profile read; pending none.")]),
            response([text("done")]),
        ]
    )
    world.chat(client, "hello there")
    return client


@scenario(
    "context_limit_1m_clear_edits",
    "A thread at effort high and the 1M context limit: clear_tool_uses and compact triggers scale with the limit; a response reporting an applied clear_tool_uses edit is metered.",
    [
        "tests/test_agent_context.py::test_thresholds_follow_the_thread_context_limit",
        "tests/test_agent_context.py::test_usage_ledger_records_server_clears",
    ],
)
def context_limit_1m(world):
    thread_id = service.create_thread(world.conn, CHAT, now=clock(NOW))
    threads.set_options(world.conn, thread_id, threads.HIGH, 1_000_000)
    client = Recorder([response([text("hi")], context_management=cleared(30_000, 4))])
    world.chat(client, "hello", thread_id=thread_id)
    return client


def effort_scenario(effort):
    @scenario(
        f"effort_{effort}",
        f"output_config.effort = {effort}; max_tokens follows the effort ({loop.max_tokens_for(effort)}).",
        ["tests/test_agent_loop.py::test_request_shape_pins_effort_caching_and_stamp"],
    )
    def run(world):
        thread_id = service.create_thread(world.conn, CHAT, now=clock(NOW))
        threads.set_options(world.conn, thread_id, effort, threads.DEFAULT_CONTEXT_LIMIT)
        client = Recorder([response([text("ok")])])
        world.chat(client, "hello", thread_id=thread_id)
        return client

    return run


for _effort in (threads.LOW, threads.XHIGH, threads.MAX):
    effort_scenario(_effort)


@scenario(
    "subagent_research",
    "A chat run calls research(role=web): the child sub run sends its own request (allowlisted tools + web_search, digest rules and the agent prompt as the last cached system block), then the parent continues with the digest.",
    [
        "tests/test_agent_subagents.py::test_research_creates_child_thread_and_returns_digest",
        "tests/test_agent_subagents.py::test_observability_and_web_search_allowlist",
    ],
)
def subagent_research(world):
    client = Recorder(
        [
            response([tool_use("toolu_01", "research", role="web", task="Find the 2027 carnival dates")], "tool_use"),
            response([text("Findings:\n- Carnival 2027 is 8-9 February [https://carnival.test/2027]\nUncertainty: none.")]),
            response([text("Carnival 2027 is on 8-9 February.")]),
        ]
    )
    world.chat(client, "Ask the web agent when carnival 2027 is.")
    return client


@scenario(
    "project_thread",
    "A thread in a project: the project block (instructions + wiki hub) follows the cached prefix uncached; effort and context limit come from the project.",
    [
        "tests/test_projects.py::test_instructions_reach_the_system_prompt_only_for_project_threads",
        "tests/test_projects.py::test_hub_instructions_name_the_page_and_the_capture_skill",
    ],
)
def project_thread(world):
    project_id = projects.create(world.conn, NOW, name="Garden", instructions="Answer in metric units.", effort=threads.HIGH, context_limit=200_000, wiki_hub_id=HUB)
    client = Recorder([response([text("Plant 30 cm apart.")])])
    world.chat(client, "How far apart should tomatoes go?", project_id=project_id)
    return client


def docx_bytes():
    document = docx.Document()
    for paragraph in ("Hello from Word", "Second paragraph"):
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "qty", "7"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx_bytes():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(["month", "total"])
    sheet.append(["Jan", 120])
    book.create_sheet("Notes").append(["remember, commas"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@scenario(
    "chat_attachments",
    "Chat attachments without the code sandbox: PDF as a base64 document, PNG as a base64 image, Markdown, .docx and .xlsx as extracted-text documents (title, untrusted context, citations), ahead of the operator text; a second turn rebuilds them from storage.",
    [
        "tests/test_files.py::test_each_type_becomes_the_right_content_block",
        "tests/test_files.py::test_attachments_persist_across_turns_and_are_rebuilt_each_request",
        "tests/test_files.py::test_sandbox_off_sends_no_code_execution_tool_and_touches_no_files_api",
    ],
)
def chat_attachments(world):
    uploads = [
        ("report.pdf", PDF, "application/pdf"),
        ("pic.png", PNG, "image/png"),
        ("notes.md", b"# hi\nsee https://evil.test/x", "text/markdown"),
        ("brief.docx", docx_bytes(), "application/octet-stream"),
        ("sales.xlsx", xlsx_bytes(), "application/octet-stream"),
    ]
    client = Recorder([response([text("Read them.")]), response([text("Still there.")])])
    thread_id = world.chat_with_files(client, "summarize", uploads)
    world.chat_with_files(client, "and again", [("later.txt", b"one more line", "text/plain")], thread_id=thread_id, now=clock(LATER))
    return client


@scenario(
    "project_files",
    "A project file joins the first user turn of every request in the project's threads, ahead of a 'Project files' note and the operator text.",
    ["tests/test_files.py::test_project_files_are_attached_to_every_request_of_its_threads"],
)
def project_files(world):
    project_id = projects.create(world.conn, NOW, name="Taxes", instructions="", effort=threads.MEDIUM, context_limit=90_000)
    files.ingest_uploads(world.cfg, world.conn, [("rules.md", b"# Rules\nfile early", "text/markdown")], NOW, project_id=project_id)
    client = Recorder([response([text("Sure.")]), response([text("More.")])])
    thread_id = world.chat(client, "what should I do?", project_id=project_id)
    world.chat(client, "next", thread_id=thread_id, now=clock(LATER))
    return client


@scenario(
    "cron_run",
    "A scheduled job run by the scheduler: kind cron, effort high, operator-written prompt, 5m TTL, web_search offered.",
    [
        "tests/test_agent_scheduler.py::test_on_time_run_is_ok_and_uses_cron_thread",
        "tests/test_agent_loop.py::test_cron_runs_at_high_effort_without_web_search_for_nightly",
    ],
)
def cron_run(world):
    job_id = scheduler.create_job(world.conn, name="brief", prompt="Check the inbox and summarise anything urgent.", schedule="daily 12:30", tz=TZ, now=NOW)
    client = Recorder([response([text("Nothing urgent.")])])
    world.job(client, CRON, scheduler.get_job(world.conn, job_id)["next_due"])
    return client


@scenario(
    "nightly_run",
    "The seeded nightly job: kind nightly, effort high, NIGHTLY_PROMPT with the wiki-capture skill expanded, nightly tool set, no web_search.",
    [
        "tests/test_agent_scheduler.py::test_nightly_is_seeded_and_runs_as_nightly",
        "tests/test_agent_scheduler.py::test_nightly_tool_set_is_exact",
    ],
)
def nightly_run(world):
    scheduler.ensure_nightly(world.conn, TZ, NOW)
    job = world.conn.execute("SELECT * FROM jobs WHERE kind = ?", (NIGHTLY,)).fetchone()
    client = Recorder([response([text("Nothing changed.")])])
    world.job(client, NIGHTLY, job["next_due"])
    return client


def run_scenario(fn):
    with tempfile.TemporaryDirectory(prefix="atrium-golden-") as tmp:
        world = World(Path(tmp))
        try:
            client = fn(world)
            assert not client.responses, f"{len(client.responses)} scripted responses unused"
            runs, usage_rows = world.records()
        finally:
            world.conn.close()
    failed = [r for r in runs if r["state"] != "done"]
    assert not failed, failed
    return client.exchanges, runs, usage_rows


def cached_prefix(body):
    system = body["system"]
    last = max(i for i, b in enumerate(system) if "cache_control" in b)
    return {"system": system[: last + 1], "tools": body.get("tools", [])}


def first_request(world, kind, message, now, **thread):
    client = Recorder([response([text("ok")])])
    if kind == SUB:
        agent = subagents.get(world.conn, thread.pop("role"))
        parent = service.create_thread(world.conn, CHAT, now=now)
        child = subagents.create_child(world.conn, parent, None, agent, f"{agent['name']}: {message}", now)
        service.run_once(world.cfg, world.conn, child, message, kind=SUB, client=client, env={}, now=now, system_extra=subagents.system_addendum(agent), allowed_tools=agent["tools"])
    else:
        world.chat(client, message, kind=kind, now=now, **thread)
    return client.exchanges[0]["request"]["body"]


def stability_kinds(world):
    """(label, kind, thread options, setup) for every run kind and prefix variant the loop builds."""
    kinds = [("chat", CHAT, {}, None), ("chat_ttl_1h", CHAT, {}, lambda w: budget.set_setting(w.conn, loop.CACHE_TTL_KEY, "1h"))]
    kinds.append(("chat_project", CHAT, {"project": True}, None))
    kinds += [("cron", CRON, {}, None), ("nightly", NIGHTLY, {}, None)]
    kinds += [(f"sub:{a['name']}", SUB, {"role": a["name"]}, None) for a in subagents.enabled(world.conn)]
    return kinds


def labels_of(world, system, extra, project_text):
    known = {loop.SYSTEM: "system", profile_block(world.conn): "profile", skills.index_block(world.conn): "skills_index"}
    if extra:
        known[extra] = "role_addendum"
    if project_text:
        known[project_text] = "project"
    return [known[b["text"]] for b in system]


def prompt_stability():
    """Builds each kind's first request twice, in fresh instances at different dates, thread ids and operator
    text, and requires the cached prefix (system through its last breakpoint, all tools) to be identical."""
    texts, definitions, kinds = {}, {}, {}
    with tempfile.TemporaryDirectory(prefix="atrium-golden-") as tmp:
        probe = World(Path(tmp) / "probe")
        labels = stability_kinds(probe)
        probe.conn.close()
        for label, kind, options, setup in labels:
            bodies = []
            for i, (start, message) in enumerate(((NOW, "hello"), (LATER, "a different question about https://example.test"))):
                world = World(Path(tmp) / f"{label.replace(':', '-')}-{i}")
                try:
                    if setup:
                        setup(world)
                    for _ in range(i * 3):
                        service.create_thread(world.conn, CHAT, now=clock(start))
                    thread = dict(options)
                    project_text = None
                    if thread.pop("project", False):
                        thread["project_id"] = projects.create(world.conn, start, name="Garden", instructions="Answer in metric units.", effort=threads.HIGH, context_limit=200_000, wiki_hub_id=HUB)
                    body = first_request(world, kind, message, clock(start), **thread)
                    if "project_id" in thread:
                        project_text = projects.system_block(world.conn, {"project_id": thread["project_id"]})
                    extra = subagents.system_addendum(subagents.get(world.conn, options["role"])) if kind == SUB else None
                    bodies.append((body, labels_of(world, body["system"], extra, project_text)))
                finally:
                    world.conn.close()
            (first, names), (second, _) = bodies
            prefix = cached_prefix(first)
            assert prefix == cached_prefix(second), f"{label}: cached prefix differs between instances"
            assert "Current date and time" not in dumps(prefix), f"{label}: date stamp in the cached prefix"
            assert all(b["type"] == "text" and "Current date and time" not in b["text"] for b in first["system"]), label
            for block, name in zip(first["system"], names, strict=True):
                assert texts.setdefault(name if name != "role_addendum" else f"role_addendum:{options['role']}", block["text"]) == block["text"], name
            for tool in first.get("tools", []):
                bare = {k: v for k, v in tool.items() if k != "cache_control"}
                assert definitions.setdefault(tool["name"], bare) == bare, f"{tool['name']} differs between kinds"
            kinds[label] = {
                "kind": kind,
                "system": [
                    {"block": name if name != "role_addendum" else f"role_addendum:{options['role']}", "cache_control": b.get("cache_control")}
                    for b, name in zip(first["system"], names, strict=True)
                ],
                "cached_system_blocks": len(prefix["system"]),
                "tools": [t["name"] for t in first.get("tools", [])],
                "tool_cache_control": [{"name": t["name"], "cache_control": t["cache_control"]} for t in first.get("tools", []) if "cache_control" in t],
                "history_cache_control": first["messages"][-1]["content"][-1]["cache_control"],
            }
    return {"freeze_commit": FREEZE_COMMIT, "kinds": kinds, "system_texts": texts, "tool_definitions": definitions}


README = """# Anthropic request goldens

Generated by `tools/golden/gen_anthropic.py` (`just golden`) from the frozen Python app at commit
`{commit}`. Do not edit by hand; after each post-freeze Python fix, regenerate from main and commit the result; the provenance is the git history of testdata.

Each `<scenario>.json` holds `exchanges`: the HTTP request the Python anthropic SDK (version {sdk}) sends
(`method`, `path`, `query`, the `anthropic-beta`/`anthropic-version`/`content-type` headers and the parsed JSON
`body`) and the scripted response the loop received (tests/fake_anthropic.py shape). `runs` and `usage` are the
rows the run left in the `runs` and `usage` tables. Compare bodies with canonical-JSON equality
(`internal/testutil/golden`), never byte equality: key order, HTML escaping and number formatting differ
between the SDKs.

`prompt_stability.json` lists, per run kind, the system block order with each block's cache_control, the tool
order and its breakpoint, and the full text of every system block and tool definition. The generator builds every
kind twice in separate instances at different dates, thread ids and operator text and fails unless the cached
prefix (system through its last breakpoint, plus all tools) is identical and free of the date stamp.

Excluded (not ported in Phase A): {excluded}.

Known divergences for the Go port (also in `manifest.json` under `divergences`):
{divergences}

`manifest.json` lists every scenario with its description and the Python tests it mirrors.

| Scenario | Requests | Mirrors |
| --- | --- | --- |
{rows}
"""


def main():
    target = out_dir(SLICE)
    if target.exists():
        shutil.rmtree(target)
    target.mkdir(parents=True)
    registry.all_tools = ported_tools
    manifest = []
    for name, description, sources, fn in SCENARIOS:
        exchanges, runs, usage_rows = run_scenario(fn)
        write_json(
            target / f"{name}.json",
            {"scenario": name, "description": description, "source_tests": sources, "freeze_commit": FREEZE_COMMIT, "exchanges": exchanges, "runs": runs, "usage": usage_rows},
        )
        manifest.append({"name": name, "file": f"{name}.json", "description": description, "source_tests": sources, "requests": len(exchanges)})
    write_json(target / "prompt_stability.json", prompt_stability())
    write_json(
        target / "manifest.json",
        {
            "slice": SLICE,
            "freeze_commit": FREEZE_COMMIT,
            "generator": str(Path(__file__).resolve().relative_to(ROOT)),
            "sdk": f"anthropic-python {anthropic.__version__}",
            "timezone": TZ,
            "excluded": EXCLUDED,
            "divergences": DIVERGENCES,
            "scenarios": manifest,
            "prompt_stability": "prompt_stability.json",
        },
    )
    rows = "\n".join(f"| `{m['name']}` | {m['requests']} | {', '.join(f'`{s}`' for s in m['source_tests'])} |" for m in manifest)
    excluded = "; ".join([*(f"tool `{t}`" for t in EXCLUDED["tools"]), *EXCLUDED["paths"]])
    divergences = "\n".join(f"- {d['where']}: {d['what']}; {d['owner']}." for d in DIVERGENCES)
    (target / "README.md").write_text(README.format(commit=FREEZE_COMMIT, sdk=anthropic.__version__, excluded=excluded, divergences=divergences, rows=rows), encoding="utf-8")


if __name__ == "__main__":
    main()
